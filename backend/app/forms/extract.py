"""Extracts plain text from an uploaded document so it can be fed into the
same form-generation prompt as a typed description. Shared by Forms'
document-to-form generator and Chat with Document - both just call
extract_text() and don't care about the underlying file format."""
import io

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx")


def extract_text(filename: str, content: bytes) -> str:
    lower = filename.lower()

    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip()

    if lower.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        return "\n".join(paragraphs).strip()

    if lower.endswith(".pptx"):
        presentation = Presentation(io.BytesIO(content))
        chunks = []
        for slide_num, slide in enumerate(presentation.slides, start=1):
            slide_lines = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            if slide_lines:
                chunks.append(f"Slide {slide_num}:\n" + "\n".join(slide_lines))
        return "\n\n".join(chunks).strip()

    raise ValueError(f"Unsupported file type: {filename}")
